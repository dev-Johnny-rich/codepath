from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

names = [
    "GB18030 Bitmap",
    "NISC18030",
    "GB18030Bitmap",
    "STHeitiSC-Medium",
    "Heiti SC",
    "Heiti SC Medium",
    "Hiragino Sans GB W3",
    "HiraginoSansGB-W3",
    "AppleGothic",
]
d = Document()
for name in names:
    p = d.add_paragraph()
    r = p.add_run(name + "  中文字体测试 项目方案 错因诊断")
    r.font.name = name
    r.font.size = Pt(18)
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
d.save("/private/tmp/font-test.docx")
